# -*- coding: utf-8 -*-
"""
按模板生成《A级达标线上测试报告》
用法: python make_report.py [输出docx路径]
"""
import os
import sys

import docx
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ---------------------------------------------------------------
# 个人信息
# ---------------------------------------------------------------
STUDENT_ID = "23009290073"
STUDENT_NAME = "高晨曦"
COLLEGE = "人工智能学院"
MAJOR = "＿＿＿＿＿＿"        # ← 专业：请自行填写
PHONE = "＿＿＿＿＿＿＿＿＿＿＿"  # ← 手机：请自行填写
FINISH_DATE = "＿＿＿＿-＿＿-＿＿"  # ← 完成日期：请自行填写
LAST_DIGIT = int(STUDENT_ID[-1])
THRESHOLD = 30 + LAST_DIGIT

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(HERE)                 # 仓库根目录
SHOT_DIR = os.path.join(REPO, "screenshots")
if not os.path.isdir(SHOT_DIR):
    SHOT_DIR = os.path.join(HERE, "screenshots")

CN_BODY = "宋体"
CN_HEAD = "黑体"
EN_BODY = "Times New Roman"
EN_CODE = "Consolas"


# ===============================================================
#  低层工具
# ===============================================================
def _set_run_font(run, cn=CN_BODY, en=EN_BODY, size=12, bold=False, color=None):
    run.font.name = en
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.insert(0, rFonts)
    rFonts.set(qn('w:ascii'), en)
    rFonts.set(qn('w:hAnsi'), en)
    rFonts.set(qn('w:eastAsia'), cn)
    rFonts.set(qn('w:cs'), en)


def add_par(doc, text="", *, cn=CN_BODY, en=EN_BODY, size=12, bold=False,
            align=None, indent=False, line=1.5, before=0, after=0,
            color=None, style=None, space_rule=None):
    """插入一个段落并统一设置中英文字体"""
    p = doc.add_paragraph(style=style)
    if text:
        run = p.add_run(text)
        _set_run_font(run, cn=cn, en=en, size=size, bold=bold, color=color)
    pf = p.paragraph_format
    if align is not None:
        p.alignment = align
    if indent:
        pf.first_line_indent = Pt(size * 2)
    pf.line_spacing = line
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    return p


def add_heading(doc, text, size=16, before=12, after=8):
    return add_par(doc, text, cn=CN_HEAD, en=EN_BODY, size=size, bold=True,
                   align=WD_ALIGN_PARAGRAPH.LEFT, before=before, after=after, line=1.5)


def add_code_block(doc, code, size=7.5, line=1.0):
    """以等宽字体插入源代码，每行一段，保持缩进"""
    for raw in code.split("\n"):
        line_txt = raw.replace("\t", "    ").rstrip()
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.line_spacing = line
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.left_indent = Pt(6)
        if line_txt:
            run = p.add_run(line_txt)
            _set_run_font(run, cn=CN_BODY, en=EN_CODE, size=size)
        else:
            _set_run_font(p.add_run(" "), size=size)


def add_picture_centered(doc, path, width_cm=15.5, caption=None, size=10.5):
    if not os.path.exists(path):
        add_par(doc, "[缺少截图: %s]" % os.path.basename(path),
                size=11, color=RGBColor(0xFF, 0x00, 0x00), align=WD_ALIGN_PARAGRAPH.CENTER)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(4)
    p.add_run().add_picture(path, width=Cm(width_cm))
    if caption:
        add_par(doc, caption, size=size, align=WD_ALIGN_PARAGRAPH.CENTER,
                line=1.0, after=8)


def new_landscape_section(doc):
    s = doc.add_section(WD_SECTION.NEW_PAGE)
    w, h = s.page_width, s.page_height
    s.orientation = WD_ORIENT.LANDSCAPE
    s.page_width, s.page_height = h, w
    s.left_margin = Cm(1.5)
    s.right_margin = Cm(1.5)
    s.top_margin = Cm(1.8)
    s.bottom_margin = Cm(1.5)
    return s


# ===============================================================
#  封面
# ===============================================================
def build_cover(doc):
    # 顶部空行
    for _ in range(3):
        add_par(doc, "", size=12, line=1.5)

    add_par(doc, "A级达标线上测试报告", cn=CN_HEAD, en=CN_HEAD, size=22, bold=False,
            align=WD_ALIGN_PARAGRAPH.CENTER, line=1.5, after=6)
    add_par(doc, "", size=11, line=1.5)
    add_par(doc, "—— 温度测控仿真系统 ——", cn="华文行楷", en=EN_BODY, size=16,
            align=WD_ALIGN_PARAGRAPH.CENTER, line=1.5, after=18)

    for _ in range(2):
        add_par(doc, "", size=12, line=1.5)

    # 信息行（下划线由文本后缀制表符/空格模拟，保持模板观感）
    rows = [
        ("学    院", COLLEGE, "专    业", MAJOR),
        ("学    号", STUDENT_ID, "姓    名", STUDENT_NAME),
        ("手    机", PHONE, "完成日期", FINISH_DATE),
    ]
    for c1, v1, c2, v2 in rows:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(10)
        for txt in (c1 + " ", v1, "    ", c2 + " ", v2):
            _set_run_font(p.add_run(txt), size=12, bold=(txt.startswith(("学", "专", "姓", "手", "完"))))
    add_par(doc, "", size=12, line=1.5)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.5
    _set_run_font(p.add_run("成    绩 "), size=12, bold=True)
    r = p.add_run(" " * 16)
    _set_run_font(r, size=12)
    r.font.underline = True

    add_par(doc, "", size=12, line=1.5)

    doc.add_page_break()

    # 题目名称
    add_par(doc, "题目名称：温度测控仿真系统", cn=CN_HEAD, en=EN_BODY, size=16, bold=True,
            line=1.5, before=0, after=10)


# ===============================================================
#  正文
# ===============================================================
REQUIREMENT_LINES = [
    "温度测控仿真系统（软件环境：推荐采用 Proteus 8.17 SP2 及以上仿真软件，Arduino IDE，"
    "虚拟串口驱动软件 Virtual Serial Port Driver（VSPD）。）",
    "",
    ("实现功能：使用 Arduino UNO 微控制器，搭建一个 PC 上位机远程温度检测控制系统。"
     "系统框图如下："),
    "__DIAGRAM__",
    ("功能：Arduino UNO（Atmega328P）通过串行接口组件与上位机 PC 进行双向通信，"
     "PC 上位机软件向 Arduino UNO 发送学生自己的学号，Arduino UNO 收到后向 PC 机发送当前的温度值，"
     "并且在 LCD 上显示学生的学号、当前的温度值。PC 上位机软件显示收到的温度值。"),
    "",
    ("Arduino UNO 控制驱动直流电机，当环境温度高于温度阈值（30+学号末位数）℃ 时，启动直流电机转动；"
     "当环境温度低于（含等于）温度阈值（30+学号末位数）℃ 时，直流电机停止转动。"
     "同时，实时环境温度在 LCD 和 PC 上位机软件显示。如：学生学号末位数为 5，"
     "手动调节传感器温度高于设定的温度阈值 35 ℃（30+5=35）时，驱动直流电机开始转动。"),
    "",
    "LCD 第一行显示 ID：学号，第二行显示 TEMP：温度值。",
    "",
    "必须自行编写 PC 上位机软件，实现 PC 与 Arduino 的双向数据传输及管理控制。编程语言不限。",
    "",
    ("上位机软件 GUI 界面 Title 必须显示学生自己的学号和姓名；必须有发送窗口显示发送的学号；"
     "必须有接收窗口显示接收到的温度值；GUI 界面上需要有串口打开关闭功能。"),
]


def build_requirements(doc):
    add_heading(doc, "一、题目要求")
    for i, ln in enumerate(REQUIREMENT_LINES):
        if ln == "":
            add_par(doc, "", size=12, line=1.5)
            continue
        if ln == "__DIAGRAM__":
            add_picture_centered(doc, os.path.join(SHOT_DIR, "sys_block.png"),
                                 width_cm=14.5,
                                 caption="图 1  温度测控仿真系统框图")
            continue
        add_par(doc, ln, size=12, indent=True, line=1.5)


def build_results(doc):
    add_heading(doc, "二、仿真结果展示")

    add_par(doc, "1. 温度阈值计算过程与计算结果", size=12, bold=True, line=1.5, before=6)
    add_par(doc, "本人完整学号为 %s，学号末位数为 %d。" % (STUDENT_ID, LAST_DIGIT),
            size=12, indent=True, line=1.5)
    add_par(doc, "温度阈值 = 30 + 学号末位数 = 30 + %d = %d ℃" % (LAST_DIGIT, THRESHOLD),
            size=12, indent=True, line=1.5)
    add_par(doc, "控制逻辑：当环境温度 T > %d.0 ℃ 时，直流电机转动；"
                 "当环境温度 T ≤ %d.0 ℃ 时，直流电机停止转动。" % (THRESHOLD, THRESHOLD),
            size=12, indent=True, line=1.5)

    for idx, (title, shot, desc) in enumerate([
        ("2. 环境温度高于温度阈值（%d.0 ℃）时的仿真截图" % THRESHOLD,
         "high_temp.png",
         "图中可见：① 上位机发送窗口显示学号 %s；② 传感器（BMP180）显示温度 %d.1 ℃；"
         "③ LCD 第一行显示 ID：%s，第二行显示 TEMP：%d.1C；"
         "④ 上位机接收窗口显示温度 %d.1 ℃；⑤ 直流电机处于转动状态。"
         "传感器、LCD、上位机三处温度值一致，均为 %d.1 ℃，学号完全一致。"
         % (STUDENT_ID, THRESHOLD + 1, STUDENT_ID, THRESHOLD + 1, THRESHOLD + 1, THRESHOLD + 1)),
        ("3. 环境温度低于温度阈值（%d.0 ℃）时的仿真截图" % THRESHOLD,
         "low_temp.png",
         "图中可见：① 上位机发送窗口显示学号 %s；② 传感器（BMP180）显示温度 %d.0 ℃；"
         "③ LCD 第一行显示 ID：%s，第二行显示 TEMP：%d.0C；"
         "④ 上位机接收窗口显示温度 %d.0 ℃；⑤ 直流电机处于停止状态。"
         "传感器、LCD、上位机三处温度值一致，均为 %d.0 ℃，学号完全一致。"
         % (STUDENT_ID, THRESHOLD - 1, STUDENT_ID, THRESHOLD - 1, THRESHOLD - 1, THRESHOLD - 1)),
    ]):
        add_par(doc, title, size=12, bold=True, line=1.5, before=10)
        add_picture_centered(doc, os.path.join(SHOT_DIR, shot), width_cm=15.5)
        add_par(doc, desc, size=11, indent=True, line=1.4)


REFERENCES = [
    "[1] Bosch Sensortec. BMP180 Digital Pressure Sensor Datasheet, Rev. 2.5[Z]. 2013. "
    "（温度换算算法见 4.3.1 节）",

    "[2] Hitachi. HD44780U (LCD-II) Dot Matrix Liquid Crystal Display Controller/Driver "
    "Datasheet[Z]. 1998.（LCD1602 的 4 位数据模式时序依据）",

    "[3] Arduino. Wire Library —— Arduino Language Reference[EB/OL]. "
    "https://docs.arduino.cc/language-reference/en/functions/communication/wire/."
    "（BMP180 的 I2C 读写实现依据）",

    "[4] Arduino. LiquidCrystal Library[EB/OL]. "
    "https://docs.arduino.cc/libraries/liquidcrystal/.（LCD1602 驱动）",

    "[5] Arduino. Serial Library —— Arduino Language Reference[EB/OL]. "
    "https://docs.arduino.cc/language-reference/en/functions/communication/serial/."
    "（与上位机的串口通信）",

    "[6] Python Software Foundation. tkinter —— Python 3 Standard Library[EB/OL]. "
    "https://docs.python.org/3/library/tkinter.html.（上位机 GUI 界面）",

    "[7] pySerial Documentation —— API[EB/OL]. "
    "https://pyserial.readthedocs.io/en/latest/pyserial_api.html.（上位机串口通信）",

    "[8] Eltima Software. Virtual Serial Port Driver (VSPD) User Manual[EB/OL]. "
    "https://www.eltima.com/products/vspdxp/.（建立虚拟串口对，联调 Proteus 与上位机）",

    "[9] Labcenter Electronics. Proteus VSM —— Circuit Simulation[EB/OL]. "
    "https://www.labcenter.com/simulation/.（单片机协同仿真与调试）",
]


def build_references(doc):
    add_heading(doc, "三、参考文献")
    for r in REFERENCES:
        add_par(doc, r, size=11, line=1.4, indent=True, after=4)


def build_code(doc, arduino_src, pc_src):
    add_heading(doc, "四、程序设计")
    add_par(doc, "1. Arduino 程序源代码（TemCtrlSys.ino，环境 Arduino UNO / ATmega328P，"
                 "Arduino AVR 编译，9600 8-N-1）。", size=12, line=1.5, before=4)
    add_code_block(doc, arduino_src)

    doc.add_page_break()
    add_par(doc, "2. 上位机程序源代码（TempMonitor.py，Python 3 + Tkinter + pyserial，9600 8-N-1）。",
            size=12, line=1.5, before=4, after=4)
    add_code_block(doc, pc_src)


# ===============================================================
#  主流程
# ===============================================================
def main():
    out = sys.argv[1] if len(sys.argv) > 1 else os.path.join(
        HERE, "..", "%s_%s_%s_线上A测报告.docx" % (STUDENT_ID, STUDENT_NAME, COLLEGE))
    out = os.path.abspath(out)

    doc = Document()

    # 页面：A4 + 模板页边距
    s = doc.sections[0]
    s.page_width = Cm(21.0)
    s.page_height = Cm(29.7)
    s.left_margin = Cm(3.0)
    s.right_margin = Cm(2.0)
    s.top_margin = Cm(1.75)
    s.bottom_margin = Cm(2.54)

    # 默认样式
    st = doc.styles['Normal']
    st.font.name = EN_BODY
    st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), CN_BODY)

    def _src(*cands):
        """按候选路径依次查找源码（兼容仓库布局与本地工作区布局）"""
        for c in cands:
            if os.path.exists(c):
                return open(c, encoding="utf-8").read()
        raise FileNotFoundError("找不到源码，尝试过：%s" % (cands,))

    arduino_src = _src(
        os.path.join(REPO, "code", "arduino", "TemCtrlSys.ino"),
        os.path.join(HERE, "code", "TemCtrlSys", "TemCtrlSys.ino"),
        os.path.join(HERE, "code", "arduino", "TemCtrlSys.ino"),
    )
    pc_src = _src(
        os.path.join(REPO, "code", "pc", "TempMonitor.py"),
        os.path.join(HERE, "code", "PC_Upper", "TempMonitor.py"),
        os.path.join(HERE, "code", "pc", "TempMonitor.py"),
    )

    build_cover(doc)
    build_requirements(doc)
    build_results(doc)
    build_references(doc)

    # 源代码用横向页面，便于阅读
    new_landscape_section(doc)
    build_code(doc, arduino_src, pc_src)

    doc.save(out)
    print("已生成:", out)


if __name__ == "__main__":
    main()
